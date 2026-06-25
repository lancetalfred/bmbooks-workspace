"""
Converts BMBooks Product Hub markdown files into a single Excel workbook.
One sheet per section. Tables rendered as styled Excel tables.
"""

import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Paths
HUB_DIR = "/Users/lalfred/Downloads/ExportBlock-3301f797-e22d-4f53-b923-2304979895db-Part-1/BMBooks - Product Hub"
OUTPUT = "/Users/lalfred/Documents/claude/BMBooks_Product_Hub.xlsx"

# Brand colours
PURPLE = "430451"
LIGHT_PURPLE = "EDE0F0"
GREY = "F5F5F5"
WHITE = "FFFFFF"

FILES = [
    ("Action Items",         "3 Action Items 2e5a2823f0878025a102fdbf8ad8f5b9.md"),
    ("Roadmap",              "2 Roadmap 2baa2823f087809b9091d5bf18605e5d.md"),
    ("Future State",         "6 Future State & Strategy 2b8a2823f08780118937cd39489eb801.md"),
    ("Release Notes",        "9 Release Notes 2b8a2823f087803b85dedbf7a049be15.md"),
    ("Platform Cost",        "13 Platform Cost Comparison.md"),
]


def parse_md(path):
    """Parse a markdown file into blocks: headings, tables, and text paragraphs."""
    with open(path, encoding="utf-8") as f:
        lines = f.readlines()

    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")

        # Heading
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("# ").strip()
            blocks.append({"type": "heading", "level": level, "text": text})
            i += 1

        # Table — detect by pipe character
        elif "|" in line and i + 1 < len(lines) and "|" in lines[i + 1] and "---" in lines[i + 1]:
            table_lines = [line]
            i += 1  # skip separator
            i += 1
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i].rstrip("\n"))
                i += 1
            # Parse header and rows
            header = [c.strip() for c in table_lines[0].split("|") if c.strip()]
            rows = []
            for tl in table_lines[1:]:
                row = [c.strip() for c in tl.split("|") if c.strip()]
                if row:
                    rows.append(row)
            blocks.append({"type": "table", "header": header, "rows": rows})

        # Horizontal rule
        elif line.strip() in ("---", "***", "___"):
            blocks.append({"type": "rule"})
            i += 1

        # Non-empty text
        elif line.strip():
            blocks.append({"type": "text", "text": line.strip()})
            i += 1

        else:
            i += 1

    return blocks


def thin_border():
    side = Side(style="thin", color="CCCCCC")
    return Border(left=side, right=side, top=side, bottom=side)


def write_sheet(ws, blocks):
    row = 1

    for block in blocks:
        btype = block["type"]

        if btype == "heading":
            cell = ws.cell(row=row, column=1, value=block["text"])
            if block["level"] == 1:
                cell.font = Font(bold=True, size=14, color=WHITE)
                cell.fill = PatternFill("solid", fgColor=PURPLE)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[row].height = 22
            elif block["level"] == 2:
                cell.font = Font(bold=True, size=12, color=PURPLE)
                cell.fill = PatternFill("solid", fgColor=LIGHT_PURPLE)
                ws.row_dimensions[row].height = 18
            else:
                cell.font = Font(bold=True, size=11)
            row += 1

        elif btype == "table":
            header = block["header"]
            # Header row
            for col_idx, col_name in enumerate(header, 1):
                cell = ws.cell(row=row, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color=WHITE, size=10)
                cell.fill = PatternFill("solid", fgColor=PURPLE)
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                cell.border = thin_border()
            ws.row_dimensions[row].height = 18
            row += 1

            # Data rows
            for r_idx, data_row in enumerate(block["rows"]):
                fill = PatternFill("solid", fgColor=GREY if r_idx % 2 == 0 else WHITE)
                for col_idx, value in enumerate(data_row, 1):
                    cell = ws.cell(row=row, column=col_idx, value=value)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.fill = fill
                    cell.border = thin_border()
                    cell.font = Font(size=10)
                ws.row_dimensions[row].height = 40
                row += 1

            # Auto-width columns
            for col_idx in range(1, len(header) + 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(
                    len(str(ws.cell(r, col_idx).value or ""))
                    for r in range(row - len(block["rows"]) - 1, row)
                )
                ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 45)

            row += 1  # gap after table

        elif btype == "text":
            text = block["text"]
            # Strip leading markdown chars
            text = re.sub(r"^\*\*(.+?)\*\*$", r"\1", text)
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"^\*\s+", "• ", text)
            text = re.sub(r"^-\s+", "• ", text)
            text = re.sub(r"^\d+\.\s+", lambda m: m.group(0), text)
            cell = ws.cell(row=row, column=1, value=text)
            cell.alignment = Alignment(wrap_text=True)
            cell.font = Font(size=10)
            ws.row_dimensions[row].height = 15
            row += 1

        elif btype == "rule":
            row += 1  # just a blank row

    # Freeze top row, set first column width
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, 30)


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    for tab_name, filename in FILES:
        path = f"{HUB_DIR}/{filename}"
        try:
            blocks = parse_md(path)
        except FileNotFoundError:
            print(f"  SKIPPED (not found): {filename}")
            continue

        ws = wb.create_sheet(title=tab_name)
        ws.sheet_view.showGridLines = False
        write_sheet(ws, blocks)
        print(f"  Written: {tab_name}")

    wb.save(OUTPUT)
    print(f"\nSaved to: {OUTPUT}")


if __name__ == "__main__":
    main()
